import os
import uuid
from datetime import datetime
from sanic import Blueprint
from sanic.response import json
from parser.middleware.auth import require_auth, require_permission
from parser.models.batch import create_batch, update_batch_status, insert_log
from parser.core.pipeline import Pipeline
from parser.utils.config_loader import match_template
import openpyxl

bp = Blueprint("upload", url_prefix="/api")

UPLOAD_DIR = "uploads"


@bp.post("/upload")
@require_auth
@require_permission("data:upload")
async def upload(request):
    if "file" not in request.files:
        return json({"error": "no file"}, status=400)

    file = request.files["file"]
    project_id = int(request.form.get("project_id", 0))
    ym = request.form.get("ym", datetime.now().strftime("%Y-%m"))
    batch_no = request.form.get("batch_no", f"B{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6]}")

    if project_id <= 0:
        return json({"error": "invalid project_id"}, status=400)

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    filepath = os.path.join(UPLOAD_DIR, f"{batch_no}.xlsx")
    with open(filepath, "wb") as f:
        f.write(file.body)

    file_size = os.path.getsize(filepath)

    pool = request.app.ctx.pool
    batch_id = await create_batch(pool, batch_no=batch_no, project_id=project_id,
                                   ym=ym, uploaded_by=request.ctx.user_id,
                                   file_name=file.name, file_size=file_size)

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet_results = []
        all_success = True
        any_success = False

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            config = match_template(sheet_name)

            if not config:
                await insert_log(pool, batch_id, sheet_name, None, "skipped")
                sheet_results.append({"name": sheet_name, "template": None, "rows": 0, "status": "skipped"})
                continue

            pipeline = Pipeline(config)
            result = pipeline.run(ws, batch_id)

            await insert_log(pool, batch_id, sheet_name, result["template_id"],
                             "matched", result["total_rows"], result["success_rows"],
                             result["error_rows"])

            if result["rows"]:
                table_name = f"data_{result['template_id']}"
                await _insert_rows(pool, table_name, result["rows"])

            sheet_results.append({
                "name": sheet_name,
                "template": result["template_id"],
                "rows": result["success_rows"],
                "status": "success" if result["error_rows"] == 0 else "partial",
            })

            if result["error_rows"] > 0:
                all_success = False
            if result["success_rows"] > 0:
                any_success = True

        if all_success and any_success:
            status = "success"
        elif any_success:
            status = "partial"
        else:
            status = "failed"

        await update_batch_status(pool, batch_id, status)

    except Exception as e:
        await update_batch_status(pool, batch_id, "failed")
        return json({"batch_id": batch_id, "batch_no": batch_no, "status": "failed", "error": str(e)}, status=500)

    return json({
        "batch_id": batch_id,
        "batch_no": batch_no,
        "status": status,
        "sheets": sheet_results,
    })


async def _insert_rows(pool, table_name, rows):
    if not rows:
        return
    json_lib = __import__("json")
    sample = rows[0]
    fixed_cols = [k for k in sample.keys() if k != "monthly_data"]

    placeholders = ", ".join(["%s"] * (len(fixed_cols) + 1))
    cols_str = ", ".join(fixed_cols + ["monthly_data"])

    async with pool.acquire() as conn:
        async with conn.cursor() as cur:
            for row in rows:
                values = [row.get(c) for c in fixed_cols] + [json_lib.dumps(row.get("monthly_data", {}), ensure_ascii=False)]
                sql = f"INSERT INTO {table_name} ({cols_str}) VALUES ({placeholders})"
                await cur.execute(sql, values)
