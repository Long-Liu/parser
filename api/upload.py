import asyncio
import json as _json
import logging
import os
import uuid
from contextlib import closing
from datetime import datetime

import aiofiles
import openpyxl
from sanic import Blueprint
from sanic.response import json

from core.pipeline import Pipeline
from db.connection import execute
from db.tables import data_table_for
from middleware.auth import require_auth, require_permission
from repositories.batch import create_batch, update_batch_status, insert_log
from utils.config_loader import match_template

logger = logging.getLogger("parser.upload")
bp = Blueprint("upload", url_prefix="/api")

UPLOAD_DIR = os.path.abspath(os.environ.get("UPLOAD_DIR", "uploads"))


def _determine_status(all_success: bool, any_success: bool) -> str:
    if not any_success:
        return "skipped"
    if all_success:
        return "success"
    return "partial"


async def _process_sheet(ws, sheet_name: str, batch_id: int) -> dict:
    config = match_template(sheet_name)
    if not config:
        await insert_log(batch_id, sheet_name, None, "skipped")
        return {"name": sheet_name, "template": None, "rows": 0, "status": "skipped"}

    pipeline = Pipeline(config)
    result = pipeline.run(ws, batch_id)

    await insert_log(
        batch_id, sheet_name, result["template_id"],
        "matched", result["total_rows"], result["success_rows"], result["error_rows"],
    )

    if result["rows"]:
        await _insert_rows(result["template_id"], config.get("columns", []), result["rows"])

    return {
        "name": sheet_name,
        "template": result["template_id"],
        "rows": result["success_rows"],
        "status": "success" if result["error_rows"] == 0 else "partial",
    }


async def _insert_rows(template_id: str, columns: list[dict], rows: list[dict]):
    if not rows:
        return
    dtable = data_table_for(template_id, columns)
    value_dicts = [
        {**{c: row.get(c) for c in row if c != "monthly_data"},
         "monthly_data": _json.dumps(row.get("monthly_data", {}), ensure_ascii=False)}
        for row in rows
    ]
    await execute(dtable.insert(), value_dicts)


@bp.post("/upload")
@require_auth
@require_permission("data:upload")
async def upload(request):
    if "file" not in request.files:
        return json({"error": "no file"}, status=400)

    file = request.files.get("file")
    if isinstance(file, list):
        file = file[0]
    try:
        project_id = int(request.form.get("project_id", 0))
    except (ValueError, TypeError):
        return json({"error": "project_id must be an integer"}, status=400)

    if project_id <= 0:
        return json({"error": "invalid project_id"}, status=400)

    ym = request.form.get("ym", datetime.now().strftime("%Y-%m"))
    batch_no = f"B{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6]}"

    os.makedirs(UPLOAD_DIR, exist_ok=True)
    filepath = os.path.join(UPLOAD_DIR, f"{batch_no}.xlsx")

    async with aiofiles.open(filepath, "wb") as f:
        await f.write(file.body)

    user_id = getattr(request.ctx, "user_id", None)
    if user_id is None:
        return json({"error": "not authenticated"}, status=401)

    batch_id = await create_batch(
        batch_no=batch_no, project_id=project_id, ym=ym,
        uploaded_by=user_id, file_name=file.name,
        file_size=os.path.getsize(filepath),
    )

    try:
        wb = await asyncio.to_thread(openpyxl.load_workbook, filepath, data_only=True)
        with closing(wb):
            sheet_results = []
            all_success = True
            any_success = False

            for sheet_name in wb.sheetnames:
                r = await _process_sheet(wb[sheet_name], sheet_name, batch_id)
                sheet_results.append(r)
                if r["status"] == "partial":
                    all_success = False
                if r["rows"] > 0:
                    any_success = True

            status = _determine_status(all_success, any_success)
            await update_batch_status(batch_id, status)

    except asyncio.CancelledError:
        raise
    except Exception:
        logger.exception("upload failed for batch %s", batch_no)
        await update_batch_status(batch_id, "failed")
        return json({
            "batch_id": batch_id, "batch_no": batch_no,
            "status": "failed", "error": "upload processing failed",
        }, status=500)
    finally:
        try:
            os.remove(filepath)
        except OSError:
            pass

    return json({
        "batch_id": batch_id, "batch_no": batch_no,
        "status": status, "sheets": sheet_results,
    })
